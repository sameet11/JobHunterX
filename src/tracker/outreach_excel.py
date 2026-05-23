"""Parse linkedin_targets.md and append rows to the outreach Excel tracker.

One row per recruiter. Columns:
  Outreach Date | Company | Job Title | Job Link |
  Recruiter Name | LinkedIn Profile | Connection Note |
  Status | Notes

Status and Notes are blank — fill them in manually as you work through the list.
New runs append rows so all history stays in one file.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

_HEADERS = [
    "Outreach Date",
    "Company",
    "Job Title",
    "Job Link",
    "Recruiter Name",
    "LinkedIn Profile",
    "Connection Note",
    "Status",
    "Notes",
]

_DATE_RE = re.compile(r"#\s+LinkedIn Outreach Targets\s+[—–-]+\s+(.+)")
_JOB_RE = re.compile(r"^##\s+(.+?)\s+[—–-]+\s+(.+)")
_LINK_RE = re.compile(r"\*\*Job:\*\*\s+(.+)")
_PERSON_RE = re.compile(r"^###\s+(.+?)\s+[—–-]+")
_PROFILE_RE = re.compile(r"\*\*Profile:\*\*\s+(https?://\S+)")
_NOTE_MARKER = "Connection-request note"


def _parse_md(text: str) -> list[dict]:
    """Return list of row dicts from linkedin_targets.md content."""
    rows: list[dict] = []

    outreach_date = ""
    m = _DATE_RE.search(text)
    if m:
        try:
            outreach_date = datetime.strptime(m.group(1).strip(), "%Y-%m-%d %H:%M").strftime("%Y-%m-%d %H:%M")
        except ValueError:
            outreach_date = m.group(1).strip()

    # Split into per-company sections on the --- dividers
    sections = re.split(r"\n---\n", text)

    for section in sections:
        lines = section.strip().splitlines()
        if not lines:
            continue

        company = job_title = job_link = ""

        for line in lines:
            m = _JOB_RE.match(line.strip())
            if m:
                company, job_title = m.group(1).strip(), m.group(2).strip()
                continue
            m = _LINK_RE.search(line)
            if m:
                job_link = m.group(1).strip()
                continue

        if not company:
            continue

        # Split into per-person sub-sections
        person_blocks = re.split(r"(?=^###\s)", section, flags=re.MULTILINE)

        for block in person_blocks:
            pm = _PERSON_RE.match(block.strip())
            if not pm:
                continue
            recruiter_name = pm.group(1).strip()

            profile_url = ""
            pm2 = _PROFILE_RE.search(block)
            if pm2:
                profile_url = pm2.group(1).strip()

            # Extract connection note: the single `> ...` line right after
            # the "Connection-request note" marker
            connection_note = ""
            blines = block.splitlines()
            for i, ln in enumerate(blines):
                if _NOTE_MARKER in ln:
                    # Next non-empty line starting with `>`
                    for nxt in blines[i + 1:]:
                        stripped = nxt.strip()
                        if stripped.startswith(">"):
                            connection_note = stripped.lstrip("> ").strip()
                            break
                    break

            rows.append({
                "Outreach Date": outreach_date,
                "Company": company,
                "Job Title": job_title,
                "Job Link": job_link,
                "Recruiter Name": recruiter_name,
                "LinkedIn Profile": profile_url,
                "Connection Note": connection_note,
                "Status": "",
                "Notes": "",
            })

    return rows


def _col_widths() -> dict[str, int]:
    return {
        "Outreach Date": 18,
        "Company": 22,
        "Job Title": 32,
        "Job Link": 45,
        "Recruiter Name": 24,
        "LinkedIn Profile": 45,
        "Connection Note": 60,
        "Status": 14,
        "Notes": 30,
    }


def _style_header_row(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)


def _apply_widths(ws, widths: dict[str, int]) -> None:
    for i, col in enumerate(_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 20)


def _make_hyperlink(ws, row: int, col: int, url: str, display: str) -> None:
    cell = ws.cell(row=row, column=col, value=display or url)
    if url:
        cell.hyperlink = url
        cell.font = Font(color="0563C1", underline="single")


def write_outreach_excel(md_path: Path, xlsx_path: Path) -> int:
    """Parse *md_path*, append new rows to *xlsx_path*, return rows added."""
    if not HAS_OPENPYXL:
        return 0

    text = md_path.read_text(encoding="utf-8")
    new_rows = _parse_md(text)
    if not new_rows:
        return 0

    widths = _col_widths()

    if xlsx_path.exists():
        wb = openpyxl.load_workbook(str(xlsx_path))
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LinkedIn Outreach"
        ws.append(_HEADERS)
        ws.row_dimensions[1].height = 22
        _style_header_row(ws)
        _apply_widths(ws, widths)
        ws.freeze_panes = "A2"

    # Deduplicate: skip rows already present (same date + recruiter + company)
    existing: set[tuple] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and row[1] and row[4]:  # date, company, recruiter
            existing.add((str(row[0]), str(row[1]), str(row[4])))

    added = 0
    for r in new_rows:
        key = (r["Outreach Date"], r["Company"], r["Recruiter Name"])
        if key in existing:
            continue
        next_row = ws.max_row + 1
        ws.cell(next_row, 1, r["Outreach Date"])
        ws.cell(next_row, 2, r["Company"])
        ws.cell(next_row, 3, r["Job Title"])
        _make_hyperlink(ws, next_row, 4, r["Job Link"], r["Job Link"])
        ws.cell(next_row, 5, r["Recruiter Name"])
        _make_hyperlink(ws, next_row, 6, r["LinkedIn Profile"], r["LinkedIn Profile"])
        ws.cell(next_row, 7, r["Connection Note"])
        ws.cell(next_row, 8, r["Status"])
        ws.cell(next_row, 9, r["Notes"])
        # Wrap connection note
        ws.cell(next_row, 7).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[next_row].height = 45
        added += 1
        existing.add(key)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(xlsx_path))
    return added
