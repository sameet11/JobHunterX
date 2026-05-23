"""Export jobs + outreach data from applied.sqlite to an Excel workbook.

Produces data/tracker.xlsx with two sheets:
  - Outreach   : all rows from the outreach table (who was contacted, status, channel)
  - Applications: all jobs with status != Scraped

Usage:
    python scripts/export_tracker.py
"""
from __future__ import annotations

import sqlite3
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

_REPO_ROOT = Path(__file__).resolve().parents[1]
_DB_PATH = _REPO_ROOT / "data" / "applied.sqlite"
_XLSX_PATH = _REPO_ROOT / "data" / "tracker.xlsx"
_CSV_OUTREACH = _REPO_ROOT / "data" / "outreach_export.csv"
_CSV_APPS = _REPO_ROOT / "data" / "applications_export.csv"

OUTREACH_COLS = [
    "recruiter_name",
    "recruiter_email",
    "recruiter_title",
    "recruiter_company",
    "channel",
    "style",
    "subject",
    "status",
    "sent_at",
    "error",
    "job_dedup_key",
    "created_at",
]

APPS_COLS = [
    "company",
    "title",
    "location",
    "salary",
    "status",
    "match_score",
    "source",
    "apply_url",
    "first_seen",
    "last_updated",
]


def _fetch(conn: sqlite3.Connection, query: str) -> tuple[list[str], list[tuple]]:
    conn.row_factory = sqlite3.Row
    cur = conn.execute(query)
    rows = cur.fetchall()
    cols = [d[0] for d in cur.description]
    return cols, [tuple(r) for r in rows]


def _autofit(ws) -> None:
    for col_cells in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max_len + 4, 60)


def _header_style(ws, n_cols: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1][:n_cols]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=False)


def export_xlsx(conn: sqlite3.Connection) -> Path:
    wb = openpyxl.Workbook()

    # ── Outreach sheet ──────────────────────────────────────────────
    ws_out = wb.active
    ws_out.title = "Outreach"

    cols_str = ", ".join(OUTREACH_COLS)
    _, rows = _fetch(conn, f"SELECT {cols_str} FROM outreach ORDER BY sent_at DESC")

    ws_out.append([c.replace("_", " ").title() for c in OUTREACH_COLS])
    for r in rows:
        ws_out.append(list(r))

    _header_style(ws_out, len(OUTREACH_COLS))
    _autofit(ws_out)

    # ── Applications sheet ──────────────────────────────────────────
    ws_apps = wb.create_sheet("Applications")

    cols_str = ", ".join(APPS_COLS)
    _, rows = _fetch(
        conn,
        f"SELECT {cols_str} FROM jobs WHERE status != 'Scraped' ORDER BY last_updated DESC",
    )

    ws_apps.append([c.replace("_", " ").title() for c in APPS_COLS])
    for r in rows:
        ws_apps.append(list(r))

    _header_style(ws_apps, len(APPS_COLS))
    _autofit(ws_apps)

    wb.save(_XLSX_PATH)
    return _XLSX_PATH


def export_csv(conn: sqlite3.Connection) -> list[Path]:
    import csv

    out_files = []

    # outreach CSV
    cols_str = ", ".join(OUTREACH_COLS)
    _, rows = _fetch(conn, f"SELECT {cols_str} FROM outreach ORDER BY sent_at DESC")
    with open(_CSV_OUTREACH, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(OUTREACH_COLS)
        w.writerows(rows)
    out_files.append(_CSV_OUTREACH)

    # applications CSV
    cols_str = ", ".join(APPS_COLS)
    _, rows = _fetch(
        conn,
        f"SELECT {cols_str} FROM jobs WHERE status != 'Scraped' ORDER BY last_updated DESC",
    )
    with open(_CSV_APPS, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(APPS_COLS)
        w.writerows(rows)
    out_files.append(_CSV_APPS)

    return out_files


def main() -> int:
    if not _DB_PATH.exists():
        print(f"ERROR: database not found at {_DB_PATH}", file=sys.stderr)
        print("Run `python main.py` first to populate it.", file=sys.stderr)
        return 1

    conn = sqlite3.connect(str(_DB_PATH))

    # Quick counts
    n_outreach = conn.execute("SELECT COUNT(*) FROM outreach").fetchone()[0]
    n_apps = conn.execute("SELECT COUNT(*) FROM jobs WHERE status != 'Scraped'").fetchone()[0]
    print(f"  Outreach rows : {n_outreach}")
    print(f"  Application rows: {n_apps}")

    if HAS_OPENPYXL:
        path = export_xlsx(conn)
        print(f"\n  Saved Excel tracker -> {path}")
        print("  Sheets: 'Outreach', 'Applications'")
    else:
        paths = export_csv(conn)
        print("\n  openpyxl not installed — exported as CSV instead.")
        print("  Install with:  pip install openpyxl")
        for p in paths:
            print(f"    {p}")

    conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
